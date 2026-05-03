import re
import io
from typing import List


def parse_file_bytes(file_bytes: bytes, file_name: str) -> str:
    """Parse file bytes into plain text based on file extension."""
    ext = file_name.rsplit('.', 1)[-1].lower() if '.' in file_name else ''

    if ext == 'pdf':
        return _parse_pdf(file_bytes)
    elif ext in ('docx', 'doc'):
        return _parse_docx(file_bytes)
    elif ext == 'txt':
        return _parse_txt(file_bytes)
    else:
        # Fallback: attempt UTF-8, then detect encoding
        return _parse_txt(file_bytes)


def parse_file_path(file_path: str, file_name: str) -> str:
    """Read a file from disk and parse it into plain text."""
    with open(file_path, 'rb') as f:
        return parse_file_bytes(f.read(), file_name)


def _parse_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text.strip())
        return '\n\n'.join(pages)
    except Exception as e:
        raise ValueError(f'PDF 解析失败: {e}')


def _parse_docx(data: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        paragraphs.append(text)
        return '\n'.join(paragraphs)
    except Exception as e:
        raise ValueError(f'DOCX 解析失败: {e}')


def _parse_txt(data: bytes) -> str:
    try:
        return data.decode('utf-8')
    except UnicodeDecodeError:
        try:
            import chardet
            detected = chardet.detect(data)
            encoding = detected.get('encoding') or 'utf-8'
            return data.decode(encoding, errors='replace')
        except Exception:
            return data.decode('utf-8', errors='replace')


# Keep legacy signature for backward compatibility
def parse_content(content: str, file_name: str) -> str:
    return content


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    """Split text into overlapping chunks, breaking on sentence boundaries."""
    if not text.strip():
        return []

    sentences = re.split(r'(?<=[。！？.!?\n])', text)
    chunks: List[str] = []
    current = ''

    for sentence in sentences:
        if len(current) + len(sentence) > chunk_size and current:
            chunks.append(current.strip())
            current = current[-overlap:] if len(current) > overlap else ''
        current += sentence

    if current.strip():
        chunks.append(current.strip())

    return chunks
